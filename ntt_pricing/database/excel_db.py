"""Load a supplier (Schneider) component pricing database.

Accepts either an ``.xlsx`` workbook (via openpyxl) or a ``.csv`` file so
that estimators can keep the rate card in whatever format is convenient.
Column names are matched flexibly — the loader looks for the *part number*,
*description* and *price* columns by a set of common aliases rather than
requiring an exact header layout.
"""
from __future__ import annotations

import csv
import os
from typing import Dict, List, Optional

from ..extraction.specs import parse_specification
from ..models import CatalogItem


_PART_ALIASES = ["part number", "part no", "partno", "reference", "ref",
                 "catalog", "catalogue", "sku", "code", "product code",
                 "material", "commercial ref"]
_DESC_ALIASES = ["description", "desc", "designation", "product", "name",
                 "product description", "details"]
_PRICE_ALIASES = ["price", "unit price", "list price", "net price", "cost",
                  "unit cost", "list", "amount"]
_CURR_ALIASES = ["currency", "curr", "ccy"]
_CAT_ALIASES = ["category", "family", "range", "group", "type", "product range"]
_MANUF_ALIASES = ["manufacturer", "brand", "make", "supplier", "vendor"]


class ComponentDatabase:
    """In-memory catalog with the parsed specification for each item."""

    def __init__(self, items: List[CatalogItem]):
        self.items = items
        self._by_part = {i.part_number.upper(): i for i in items}

    def __len__(self) -> int:
        return len(self.items)

    def by_part_number(self, part_number: str) -> Optional[CatalogItem]:
        return self._by_part.get(part_number.upper())

    @classmethod
    def load(cls, path: str) -> "ComponentDatabase":
        ext = os.path.splitext(path)[1].lower()
        if ext in (".xlsx", ".xlsm"):
            # a multi-sheet NTT/Schneider price book (per-family tabs) is
            # detected and parsed specially; otherwise treat as a flat sheet
            items = load_ntt_workbook(path)
            if items:
                return cls(items)
            rows = _read_xlsx(path)
        elif ext == ".csv":
            rows = _read_csv(path)
        else:
            raise ValueError(f"Unsupported pricing database format: {ext}")
        return cls(_rows_to_items(rows))


# --------------------------------------------------------------------------
# NTT / Schneider multi-sheet price book
# --------------------------------------------------------------------------
# Each product family is its own worksheet with a header row like
#   Descripion | ... | Qty | Ref. No. | Manufacture | price list | DIS | P.A.D | F.P
# and the net unit price in the F.P (Final Price) column.  A flat master tab
# (DGQ: Reference ID / Reference Description / CPQ EUR PRICE) is also read.
_NTT_REF = ["ref. no.", "ref no", "ref.", "reference id"]
_NTT_DESC = ["descripion", "description", "reference description"]
_NTT_PRICE = ["f.p", "final price", "cpq eur price", "p.a.d"]
_NTT_MANUF = ["manufacture", "manufacturer", "brand"]


def load_ntt_workbook(path: str) -> List[CatalogItem]:
    """Parse a multi-tab NTT price book into catalog items (empty if not that
    format, so the caller can fall back to a flat-sheet reader)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    items: List[CatalogItem] = []
    seen = set()
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        hidx, header = _ntt_header(rows)
        if hidx is None:
            continue
        rc = _ntt_col(header, _NTT_REF)
        dc = _ntt_col(header, _NTT_DESC)
        pc = _ntt_col(header, _NTT_PRICE)
        mc = _ntt_col(header, _NTT_MANUF)
        if rc is None or pc is None:
            continue
        if dc is None:
            dc = 0
        for r in rows[hidx + 1:]:
            if rc >= len(r) or r[rc] in (None, ""):
                continue
            part = str(r[rc]).strip()
            price = _to_float(r[pc]) if pc < len(r) else None
            if price is None or not part or part.upper() in seen:
                continue
            desc = str(r[dc]).strip() if dc < len(r) and r[dc] else ""
            manuf = (str(r[mc]).strip() if mc is not None and mc < len(r) and r[mc]
                     else "Schneider Electric")
            item = CatalogItem(part_number=part, description=desc, unit_price=price,
                               currency="EUR", manufacturer=manuf, category=sheet)
            item.spec = parse_specification(f"{desc} {sheet}")
            items.append(item)
            seen.add(part.upper())
    wb.close()
    # only treat as an NTT workbook if it yielded a substantial catalog
    return items if len(items) >= 50 else []


def _ntt_header(rows):
    for i, r in enumerate(rows[:20]):
        cells = [str(c).strip().lower() if c is not None else "" for c in r]
        has_ref = any(any(a in c for a in _NTT_REF) for c in cells)
        has_price = any(c in _NTT_PRICE for c in cells)
        if has_ref and has_price:
            return i, cells
    return None, None


def _ntt_col(header, names):
    for n in names:                       # exact label first
        for j, c in enumerate(header):
            if c == n:
                return j
    for n in names:                       # then contains
        for j, c in enumerate(header):
            if n in c:
                return j
    return None


def _read_csv(path: str) -> List[Dict[str, str]]:
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        return [dict(r) for r in reader]


def _read_xlsx(path: str) -> List[Dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
    out: List[Dict[str, str]] = []
    for row in rows_iter:
        if row is None or all(v is None for v in row):
            continue
        out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    wb.close()
    return out


def _rows_to_items(rows: List[Dict]) -> List[CatalogItem]:
    if not rows:
        return []
    header_map = _map_headers(list(rows[0].keys()))
    items: List[CatalogItem] = []
    for r in rows:
        part = _get(r, header_map, "part")
        desc = _get(r, header_map, "desc")
        price_raw = _get(r, header_map, "price")
        if part is None and desc is None:
            continue
        price = _to_float(price_raw)
        if price is None:
            continue
        item = CatalogItem(
            part_number=str(part or "").strip() or f"AUTO-{len(items)}",
            description=str(desc or "").strip(),
            unit_price=price,
            currency=str(_get(r, header_map, "curr") or "USD").strip() or "USD",
            manufacturer=str(_get(r, header_map, "manuf") or "Schneider Electric").strip(),
            category=(str(_get(r, header_map, "cat")).strip() if _get(r, header_map, "cat") else None),
        )
        item.spec = parse_specification(f"{item.description} {item.category or ''}")
        items.append(item)
    return items


def _map_headers(headers: List[str]) -> Dict[str, str]:
    """Map canonical field -> actual header name using alias lists."""
    lut = {h.lower().strip(): h for h in headers}
    mapping: Dict[str, str] = {}

    def find(aliases):
        for a in aliases:
            if a in lut:
                return lut[a]
        # partial contains match
        for key, orig in lut.items():
            if any(a in key for a in aliases):
                return orig
        return None

    mapping["part"] = find(_PART_ALIASES)
    mapping["desc"] = find(_DESC_ALIASES)
    mapping["price"] = find(_PRICE_ALIASES)
    mapping["curr"] = find(_CURR_ALIASES)
    mapping["cat"] = find(_CAT_ALIASES)
    mapping["manuf"] = find(_MANUF_ALIASES)
    return mapping


def _get(row: Dict, header_map: Dict[str, str], field: str):
    col = header_map.get(field)
    if col is None:
        return None
    return row.get(col)


def _to_float(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    s = "".join(ch for ch in s if ch.isdigit() or ch in ".-")
    try:
        return float(s) if s not in ("", "-", ".") else None
    except ValueError:
        return None
